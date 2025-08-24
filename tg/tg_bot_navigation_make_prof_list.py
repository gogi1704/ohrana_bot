from telegram import Update, InlineKeyboardMarkup, InlineKeyboardButton, ReplyKeyboardMarkup
from telegram.ext import (ContextTypes, ConversationHandler)

import resources
from tg import tg_bot_navigation
from tg.tg_keyboards import web_app_keyboard
from util_funs import get_prikaz_29_punkts_from_name, get_doctors_by_punkts, get_unique_counts_safe, get_base_doctors_or_tests, get_text_test_or_doctors


(
    ASK_MEN,
    ASK_MEN_40,
    ASK_WOMEN,
    ASK_WOMEN_40,
    SHOW_SUMMARY,
    WAIT_PROFESSION,
    WAIT_PROF_COUNT,
    PROFESSION_MENU
) = range(8)

CANCEL = "Отмена"

# === Handlers ===

async def start_make_list(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text("👨‍🦱 Сколько всего мужчин на предприятии?", reply_markup=ReplyKeyboardMarkup([[CANCEL]], resize_keyboard=True))
    return ASK_MEN

async def ask_men_40(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.message.text == CANCEL:
        return await cancel(update, context)
    if not update.message.text.isdigit():
        return await update.message.reply_text("Введите число.")
    context.user_data["men_total"] = int(update.message.text)
    await update.message.reply_text("👨‍🦰 Сколько мужчин старше 40?", reply_markup=ReplyKeyboardMarkup([[CANCEL]], resize_keyboard=True))
    return ASK_MEN_40

async def ask_women(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.message.text == CANCEL:
        return await cancel(update, context)
    if not update.message.text.isdigit():
        return await update.message.reply_text("Введите число.")
    context.user_data["men_40"] = int(update.message.text)
    await update.message.reply_text("👩‍🦱 Сколько всего женщин на предприятии?", reply_markup=ReplyKeyboardMarkup([[CANCEL]], resize_keyboard=True))
    return ASK_WOMEN

async def ask_women_40(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.message.text == CANCEL:
        return await cancel(update, context)
    if not update.message.text.isdigit():
        return await update.message.reply_text("⏳ Введите число:")
    context.user_data["women_total"] = int(update.message.text)
    await update.message.reply_text("👱‍♀️ Сколько женщин старше 40?", reply_markup=ReplyKeyboardMarkup([[CANCEL]], resize_keyboard=True))
    return ASK_WOMEN_40

async def show_summary(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.message.text == CANCEL:
        return await cancel(update, context)
    if not update.message.text.isdigit():
        return await update.message.reply_text("⏳ Введите число:")
    context.user_data["women_40"] = int(update.message.text)

    text = (
        f"👨‍🦰 Мужчин - {context.user_data['men_total']}   ({context.user_data['men_40']} старше 40 лет)\n"
        f"👩‍🦱 Женщин - {context.user_data['women_total']}   ({context.user_data['women_40']} старше 40 лет)"
    )

    keyboard = InlineKeyboardMarkup([[InlineKeyboardButton("Далее ➡️", callback_data="next")]])
    await update.message.reply_text(text, reply_markup=keyboard)
    return SHOW_SUMMARY

async def proceed_to_profession(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.callback_query.answer()
    await update.callback_query.message.delete()

    # kb = [["Выбрать из списка"], [CANCEL]]
    # markup = ReplyKeyboardMarkup(kb, resize_keyboard=True, one_time_keyboard=True)

    await update.callback_query.message.reply_text(
        "Необходимо выбрать профессию (нажав на кнопку ниже 👇), после чего укажете количество сотрудников на этой должности.",
        reply_markup= web_app_keyboard
    )
    return WAIT_PROFESSION

async def handle_webapp_profession(update: Update, context: ContextTypes.DEFAULT_TYPE):
    prof = update.message.web_app_data.data.strip()
    context.user_data["current_profession"] = prof
    await update.message.reply_text(f"{prof};\n🤔 Сколько сотрудников на этой должности?", reply_markup=ReplyKeyboardMarkup([[CANCEL]], resize_keyboard=True))
    return WAIT_PROF_COUNT

async def handle_profession_count(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.message.text == CANCEL:
        return await cancel(update, context)
    if not update.message.text.isdigit():
        return await update.message.reply_text("Введите число.")

    prof = context.user_data["current_profession"]
    count = int(update.message.text)
    context.user_data.setdefault("professions", []).append((prof, count))

    summary = (
        f"👨‍🦱 Мужчин - {context.user_data['men_total']} ({context.user_data['men_40']} старше 40 лет)\n"
        f"👩‍🦱 Женщин - {context.user_data['women_total']} ({context.user_data['women_40']} старше 40 лет)\n"
    )
    for p, c in context.user_data["professions"]:
        summary += f"{p} - {c}\n"

    keyboard = InlineKeyboardMarkup([
        [InlineKeyboardButton("Добавить профессию 🏗", callback_data="add")],
        [InlineKeyboardButton("Готово ✅", callback_data="done")]
    ])
    await update.message.reply_text(summary, reply_markup=keyboard)
    return PROFESSION_MENU

async def profession_menu_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.callback_query.answer()

    if update.callback_query.data == "add":
        await update.callback_query.message.reply_text(
            "🏗 Выберите следующую профессию:",
            reply_markup=web_app_keyboard
        )
        return WAIT_PROFESSION

    elif update.callback_query.data == "done":




        all_count = int(context.user_data['men_total']) + int(context.user_data['women_total'])
        base_doctors_list_all = get_base_doctors_or_tests(peoples_group= resources.base_doctors_all, count= all_count)
        base_doctors_list_women = get_base_doctors_or_tests(peoples_group=resources.base_doctors_women, count= int(context.user_data['women_total']))

        # base_doctors_list_men_40 =get_base_doctors_or_tests(peoples_group= resources.base_doctors_men_40 , count= int(context.user_data['men_40']))
        # base_doctors_list_women_40 = get_base_doctors_or_tests(peoples_group=resources.base_doctors_women_40, count= int(context.user_data['women_40']))

        base_tests_list_all = get_base_doctors_or_tests(peoples_group=resources.base_tests_all, count= all_count)
        base_tests_women = get_base_doctors_or_tests(resources.base_tests_women, count= int(context.user_data['women_total']))
        base_tests_list_women_40 = get_base_doctors_or_tests(peoples_group=resources.base_test_women_40, count= int(context.user_data['women_40']))
        base_tests_list_men_40 = get_base_doctors_or_tests(peoples_group=resources.base_test_men_40, count= int(context.user_data['men_40']))

        final_list = [[]]
        for prof, count in context.user_data.get("professions", []):
            punkts = get_prikaz_29_punkts_from_name(prof)
            doctors, tests = get_doctors_by_punkts(punkts)
            all_docs_and_tests = doctors * count + tests * count
            final_list.append(all_docs_and_tests)

        _,summary = get_unique_counts_safe(final_list)
        # summary_text = "Итоговый список:\n\n"
        text1 = get_text_test_or_doctors("Общие специалисты:\n", base_doctors_list_all)
        text2 = get_text_test_or_doctors("Специалисты для женщин:\n", base_doctors_list_women)
        # text3 = get_text_test_or_doctors("Специалисты для мужчин старше 40 лет:\n", base_doctors_list_men_40)
        # text4 = get_text_test_or_doctors("Специалисты для женщин старше 40 лет:\n", base_doctors_list_women_40)
        text3 = get_text_test_or_doctors("Общие обследования:\n", base_tests_list_all)
        text4 = get_text_test_or_doctors("Общие обследования для женщин:\n",base_tests_women )
        text5 = get_text_test_or_doctors("Обследования для женщин старше 40 лет:\n", base_tests_list_women_40)
        text6 = get_text_test_or_doctors("Обследования для мужчин старше 40 лет:\n", base_tests_list_men_40)
        text7 = get_text_test_or_doctors("Дополнительные обследования по вредностям:\n", summary)

        summary_text = f"""
        ♦️{text1}
        ♦️{text2}
        ♦️{text3}
        ♦️{text4}
        ♦️{text5}
        ♦️{text6}
        ♦️{text7}
        """

        # for item, count in summary.items():
        #     summary_text += f"– {item}: {count} шт.\n"

        # await update.callback_query.message.reply_text(summary_text)
        await send_styled_excel_table_to_user(update,context,base_doctors_list_all, base_doctors_list_women, base_tests_list_all, base_tests_women, base_tests_list_women_40,base_tests_list_men_40, summary)
        return ConversationHandler.END
    else:
        return None

import pandas as pd
from io import BytesIO
from telegram import Update
from telegram.ext import ContextTypes
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment

async def send_styled_excel_table_to_user(update: Update, context: ContextTypes.DEFAULT_TYPE, *dicts):
    # 1. Собираем данные из переданных словарей
    combined = {}
    for d in dicts:
        for name, qty in d.items():
            combined[name] = combined.get(name, 0) + qty

    # 2. Формируем DataFrame
    df = pd.DataFrame(
        list(combined.items()),
        columns=["Врач/обследование", "Количество"]
    )

    # 3. Пишем DataFrame в Excel в память
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Таблица")
        ws = writer.book["Таблица"]

        # 4. Стилизация заголовков
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F81BD")
        header_align = Alignment(horizontal="center", vertical="center")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # 5. Автоширина колонок
        for col in ws.columns:
            max_length = max(len(str(cell.value)) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max_length + 2


    output.seek(0)

    # 6. Отправляем файл пользователю
    await context.bot.send_document(
        chat_id=update.effective_chat.id,
        document=output,
        filename="таблица_осмотров.xlsx",
        caption="📊 Ваша таблица готова"
    )


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await tg_bot_navigation.start(update, context)
    return ConversationHandler.END


